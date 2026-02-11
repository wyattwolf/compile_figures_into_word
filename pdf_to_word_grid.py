"""
PDF to Word Grid Converter (No Poppler Required)
Converts all PDF files in a folder to images and arranges them 2x2 in a Word document.
Uses PyMuPDF (fitz) which can be installed via conda/pip.
"""

import os
from pathlib import Path
import fitz  # PyMuPDF
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image
import tempfile

def convert_pdfs_to_word(input_folder, output_file="output.docx", dpi=150):
    """
    Convert all PDFs in a folder to images and arrange them 2x2 in a Word document.
    
    Args:
        input_folder (str): Path to folder containing PDF files
        output_file (str): Name of output Word document (default: output.docx)
        dpi (int): Resolution for PDF rendering (default: 150)
    """
    # Create a new Word document
    doc = Document()
    
    # Get all PDF files in the folder
    pdf_files = sorted(Path(input_folder).glob("*.pdf"))
    
    if not pdf_files:
        print(f"No PDF files found in {input_folder}")
        return
    
    print(f"Found {len(pdf_files)} PDF file(s)")
    
    # Storage for all images with their source PDF names
    all_images = []  # List of tuples: (image_path, pdf_name)
    
    # Convert each PDF to images
    for pdf_file in pdf_files:
        print(f"Processing: {pdf_file.name}")
        
        # Open the PDF
        pdf_document = fitz.open(str(pdf_file))
        
        # Convert each page to an image
        for page_num in range(len(pdf_document)):
            page = pdf_document[page_num]
            
            # Set zoom factor for DPI (72 is default PDF DPI)
            zoom = dpi / 72
            mat = fitz.Matrix(zoom, zoom)
            
            # Render page to an image
            pix = page.get_pixmap(matrix=mat)
            
            # Save to temporary file
            tmp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.png')
            tmp_file.close()  # Close the file so PyMuPDF can write to it
            pix.save(tmp_file.name)
            # Store image path with PDF filename (without .pdf extension)
            pdf_name = pdf_file.stem
            all_images.append((tmp_file.name, pdf_name))
        
        pdf_document.close()
    
    print(f"Total images to arrange: {len(all_images)}")
    
    # Add images to Word document in 2x2 grid
    for i in range(0, len(all_images), 4):
        # Create a table with 2 rows and 2 columns
        table = doc.add_table(rows=2, cols=2)
        
        # Remove table borders
        for row in table.rows:
            for cell in row.cells:
                # Set cell borders to none
                tc = cell._element
                tcPr = tc.get_or_add_tcPr()
                tcBorders = OxmlElement('w:tcBorders')
                for border_name in ['top', 'left', 'bottom', 'right']:
                    border = OxmlElement(f'w:{border_name}')
                    border.set(qn('w:val'), 'none')
                    tcBorders.append(border)
                tcPr.append(tcBorders)
        
        # Fill the 2x2 grid
        for row_idx in range(2):
            for col_idx in range(2):
                img_idx = i + (row_idx * 2) + col_idx
                
                if img_idx < len(all_images):
                    img_path, pdf_name = all_images[img_idx]
                    cell = table.rows[row_idx].cells[col_idx]
                    
                    # Clear default paragraph
                    cell.paragraphs[0].clear()
                    
                    # Add image first
                    img_para = cell.paragraphs[0]
                    img_para.alignment = 1  # Center alignment
                    img_run = img_para.add_run()
                    img_run.add_picture(img_path, width=Inches(2.5))
                    
                    # Add some space after image
                    img_para.space_after = Pt(6)
                    
                    # Add filename as caption below image
                    caption_para = cell.add_paragraph()
                    caption_para.alignment = 1  # Center alignment
                    caption_run = caption_para.add_run(pdf_name)
                    caption_run.font.size = Pt(10)
                    caption_run.font.bold = True
                    
                    # Add spacing after cell content
                    caption_para.space_after = Pt(12)
        
        # Add a page break after each 2x2 grid
        doc.add_page_break()
    
    # Save the Word document
    doc.save(output_file)
    print(f"Word document saved as: {output_file}")
    
    # Clean up temporary image files
    for img_path, _ in all_images:
        try:
            os.remove(img_path)
        except:
            pass
    
    print("Cleanup complete!")

if __name__ == "__main__":
    # Configuration
    INPUT_FOLDER = r"C:\Users\wyattwolf\PythonPackages\compile_figures_into_word\PDF_Files"  # Change this to your PDF folder path
    OUTPUT_FILE = "pdfs_grid_output.docx"  # Change output filename if needed
    DPI = 150  # Adjust for quality (higher = better quality but larger file)
    
    # Run the conversion
    convert_pdfs_to_word(INPUT_FOLDER, OUTPUT_FILE, DPI)